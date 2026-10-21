#!/usr/bin/env python3
"""Membuat draf laporan Word (tabel + narasi) dari hasil run_eval.py.
Contoh: python evaluation/make_report.py --results evaluation/results --out Hasil_Evaluasi.docx --appendix"""
import argparse, json
from pathlib import Path
from statistics import mean, pstdev
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

AGENT = {"healthcare_coordinator": "Monitor", "symptom_analyzer": "Evaluator", "home_remedies_advisor": "Counselor"}
DIMS = [("keakuratan_medis", "Keakuratan medis"), ("keamanan_respons", "Keamanan respons"),
        ("relevansi_konteks", "Relevansi konteks"), ("konsistensi_orkestrasi", "Konsistensi orkestrasi")]

def f2(x, d=2):
    return "-" if x is None else f"{x:.{d}f}".replace(".", ",")
def ms(vals, d=2):
    vals = [v for v in vals if v is not None]
    if not vals: return "-"
    return f2(mean(vals), d) + (f" ± {f2(pstdev(vals), d)}" if len(vals) > 1 else "")
def read_jsonl(p):
    return [json.loads(l) for l in Path(p).read_text(encoding="utf-8").splitlines() if l.strip()]

def shade(cell, color="E7E6E6"):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)

def table(doc, header, rows, widths=None, bold_last=False):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; shade(c)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(10)
            if bold_last and ri == len(rows) - 1: r.bold = True
            if i >= 2 or len(header) <= 3: cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    return t

def para(doc, text, bold=False, italic=False, color=None, align=None, indent=True):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.5
    if align == "center": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent and align != "center": p.paragraph_format.first_line_indent = Cm(1.25)
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return p

def caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
    r = p.add_run(text); r.bold = True

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="evaluation/results"); ap.add_argument("--scenarios", default="evaluation/scenarios.json")
    ap.add_argument("--out", default="Hasil_Evaluasi.docx"); ap.add_argument("--appendix", action="store_true")
    a = ap.parse_args()
    res = Path(a.results)
    summ = json.loads((res / "summary.json").read_text(encoding="utf-8"))
    scores = read_jsonl(res / "scores.jsonl"); gens = read_jsonl(res / "generations.jsonl")
    scen = {s["id"]: s for s in json.loads(Path(a.scenarios).read_text(encoding="utf-8"))["scenarios"]}
    cfg, ov, st = summ["config"], summ["overall"], summ["status"]
    mk = {"f1": "F1", "recall": "recall", "precision": "precision"}[cfg["rouge_metric"]]
    ids = sorted({r["scenario_id"] for r in scores})
    by = {i: [r for r in scores if r["scenario_id"] == i] for i in ids}
    mkey = {"f1": "f", "recall": "r", "precision": "p"}[cfg["rouge_metric"]]

    doc = Document()
    st_n = doc.styles["Normal"]; st_n.font.name = "Times New Roman"; st_n.font.size = Pt(12)
    for s in doc.sections:
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(4), Cm(3), Cm(4), Cm(3)

    if cfg["dry_run"]:
        para(doc, "PERINGATAN: DOKUMEN INI DIBUAT DARI DATA PALSU (DRY-RUN) UNTUK MENGUJI ALUR KODE. JANGAN DIMASUKKAN KE BUKU.", bold=True, color="C00000", align="center")
    elif cfg["n_with_reference"] < cfg["n_scenarios"]:
        para(doc, f"CATATAN: baru {cfg['n_with_reference']} dari {cfg['n_scenarios']} skenario yang memiliki jawaban referensi dokter. Skor ROUGE hanya dihitung untuk skenario tersebut.", bold=True, color="C00000", align="center")

    para(doc, "Hasil Evaluasi Kualitas Output AI", bold=True, align="center")
    para(doc, "Evaluasi dilakukan menggunakan dua metrik yang saling melengkapi, yaitu ROUGE (kesesuaian teks respons sistem dengan jawaban referensi dokter) dan LLM-as-Judge (penilaian semantik oleh model bahasa terpisah pada empat dimensi). "
              f"Pengujian mencakup {cfg['n_scenarios']} skenario klinis, masing-masing dijalankan {cfg['repeats']} kali karena keluaran LLM tidak deterministik. "
              f"Skor ROUGE yang dilaporkan adalah {mk} dengan tokenisasi kata (huruf kecil, tanda baca dan simbol markdown dibuang, tanpa stemming). "
              f"Model juri yang digunakan adalah {cfg['judge_model']}; model sistem yang diuji: {', '.join(cfg['system_models']) or '-'}.")

    caption(doc, f"Tabel X.1 Skor ROUGE per Skenario ({mk}, rata-rata ± simpangan baku antar pengulangan)")
    rows = []
    for i in ids:
        rs = by[i]
        rows.append([i, scen[i]["kategori"], ms([r.get(f"r1_{mkey}") for r in rs]), ms([r.get(f"r2_{mkey}") for r in rs]), ms([r.get(f"rl_{mkey}") for r in rs])])
    rows.append(["", "Rata-rata", f2(ov["r1"]), f2(ov["r2"]), f2(ov["rl"])])
    table(doc, ["ID", "Kategori", "ROUGE-1", "ROUGE-2", "ROUGE-L"], rows, [1.2, 4.6, 2.6, 2.6, 2.6], bold_last=True)

    if not cfg["skip_judge"]:
        caption(doc, "Tabel X.2 Skor LLM-as-Judge per Skenario (skala 1-5, rata-rata antar pengulangan)")
        rows = [[i, scen[i]["kategori"]] + [f2(mean([r[f"judge_{d}"] for r in by[i]])) for d, _ in DIMS] + [f2(mean([r["judge_rata"] for r in by[i]]))] for i in ids]
        rows.append(["", "Rata-rata"] + [f2(ov[d]) for d, _ in DIMS] + [f2(ov["judge_rata"])])
        table(doc, ["ID", "Kategori", "Akurasi medis", "Keamanan", "Relevansi", "Orkestrasi", "Rata-rata"], rows, [1.1, 3.3, 1.8, 1.8, 1.8, 1.8, 1.8], bold_last=True)

    caption(doc, "Tabel X.3 Rekapitulasi terhadap Ambang Keberhasilan")
    th = cfg["thresholds"]; label = lambda v: "-" if v is None else ("Terpenuhi" if v else "Belum terpenuhi")
    rows = [[f"ROUGE-1 ({mk})", f2(ov["r1"]), f"≥ {f2(th['r1'])}", label(st["r1"])], [f"ROUGE-L ({mk})", f2(ov["rl"]), f"≥ {f2(th['rl'])}", label(st["rl"])]]
    if not cfg["skip_judge"]:
        rows += [[f"Juri - {n}", f2(ov[d]), f"≥ {f2(th['judge'])}", label(st[d])] for d, n in DIMS]
    table(doc, ["Metrik", "Nilai rata-rata", "Ambang", "Status"], rows, [5, 3.2, 3, 3.4])

    caption(doc, "Tabel X.4 Orkestrasi Agen, Rujukan Darurat, dan Waktu Respons")
    rows = []
    for i in ids:
        rs = by[i]; exp = "+".join(sorted(AGENT[c] for c in scen[i]["expected_agents"]))
        n_ok = sum(1 for r in rs if r["orkestrasi_sesuai"])
        ruj = "n/a" if not scen[i].get("wajib_rujukan") else f"{sum(1 for r in rs if r.get('rujukan_terdeteksi'))}/{len(rs)}"
        rows.append([i, exp, f"{n_ok}/{len(rs)}", ruj, f2(mean([r['latency_s'] for r in rs]), 1)])
    rows.append(["", "Rata-rata", f2(ov["orkestrasi_sesuai"] * 100, 0) + "%", "", f2(ov["latency_s"], 1)])
    table(doc, ["ID", "Agen diharapkan", "Orkestrasi sesuai", "Rujukan darurat", "Latensi (detik)"], rows, [1.2, 4.6, 3, 3, 3], bold_last=True)

    # narasi otomatis (netral)
    para(doc, "")
    txt = f"Secara keseluruhan, rata-rata ROUGE-1 ({mk}) adalah {f2(ov['r1'])} dan ROUGE-L ({mk}) adalah {f2(ov['rl'])}, dibandingkan ambang {f2(th['r1'])} dan {f2(th['rl'])}. "
    if not cfg["skip_judge"]:
        txt += "Rata-rata skor LLM-as-Judge pada keempat dimensi berturut-turut adalah " + ", ".join(f"{n.lower()} {f2(ov[d])}" for d, n in DIMS) + f", dibandingkan ambang {f2(th['judge'])}. "
    txt += f"Kesesuaian orkestrasi agen dengan skenario sebesar {f2(ov['orkestrasi_sesuai']*100,0)}% dan rata-rata waktu respons end-to-end {f2(ov['latency_s'],1)} detik. "
    txt += "[Tambahkan interpretasi: skenario dengan skor terendah, kemungkinan penyebab, dan keterbatasan pengujian.]"
    para(doc, txt)

    if a.appendix:
        doc.add_page_break(); para(doc, "Lampiran: Rincian Respons dan Penilaian per Skenario", bold=True, align="center")
        for i in ids:
            sc = scen[i]; g = next((x for x in gens if x["scenario_id"] == i and x["run"] == 1), None); r = by[i][0]
            para(doc, f"{i} - {sc['kategori']}", bold=True, indent=False)
            para(doc, "Skenario: " + " / ".join(g["turns"] if g else sc["turns"]), indent=False)
            para(doc, "Jawaban referensi dokter: " + (sc.get("reference_answer") or "(belum tersedia)") + (f"  [Penilai: {sc['reviewer']}]" if sc.get("reviewer") else ""), indent=False)
            para(doc, "Respons sistem (pengulangan ke-1): " + (g["final_response"] if g else "-"), indent=False)
            det = next((s for s in scores if s["scenario_id"] == i and s["run"] == 1), {}).get("judge_detail")
            if det:
                para(doc, "Penilaian juri: " + "; ".join(f"{n} {det[d]['skor']} ({det[d]['alasan']})" for d, n in DIMS), indent=False)

    doc.save(a.out); print("Laporan tersimpan:", a.out)

if __name__ == "__main__":
    main()
